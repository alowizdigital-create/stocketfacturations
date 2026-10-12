"""Comparaison avec la période précédente, tendance, panier moyen et export
Excel du rapport de marge — voir apps.sales.reports et sales:margin_report."""

import io
from datetime import timedelta
from decimal import Decimal

import openpyxl
import pytest
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import User
from apps.purchasing.tests.test_purchasing import _confirmed_sale, _product, _stock
from apps.sales import reports
from apps.sync.tests.factories import BoutiqueFactory
from apps.tenants.models import Membership

pytestmark = pytest.mark.django_db


@pytest.fixture
def boutique():
    return BoutiqueFactory()


@pytest.fixture
def logged(client, boutique):
    user = User.objects.create_user(email="g@test.com", password="x")
    Membership.objects.create(user=user, boutique=boutique, role=Membership.GERANT_BOUTIQUE, is_active=True)
    client.force_login(user)
    return client


def test_previous_period_has_same_length():
    d1, d2 = timezone.localdate().replace(day=10), timezone.localdate().replace(day=14)
    p1, p2 = reports.previous_period(d1, d2)
    assert p2 == d1 - timedelta(days=1) and (p2 - p1).days == 4


def test_sales_count_and_average_basket(boutique):
    today = timezone.localdate()
    product = _product(boutique.compte, purchase_price=Decimal("300"))
    _stock(boutique, product, 20)
    _confirmed_sale(boutique, product, 1, 1000)
    _confirmed_sale(boutique, product, 1, 2000)
    totals = reports.margin_report(boutique, today, today)["totals"]
    assert totals["sales_count"] == 2 and totals["average_basket"] == Decimal("1500")


def test_average_basket_none_without_sales(boutique):
    today = timezone.localdate()
    assert reports.margin_report(boutique, today, today)["totals"]["average_basket"] is None


def test_compare_totals_percent_and_zero_base():
    cur = {"revenue": Decimal("150"), "margin": Decimal("50"), "expenses_total": Decimal("0"),
           "net_profit": Decimal("-10"), "sales_count": 3}
    prev = {"revenue": Decimal("100"), "margin": Decimal("0"), "expenses_total": Decimal("0"),
            "net_profit": Decimal("-20"), "sales_count": 2}
    deltas = reports.compare_totals(cur, prev)
    assert deltas["revenue"] == Decimal("50.0") and deltas["sales_count"] == Decimal("50.0")
    assert deltas["margin"] is None          # base nulle : pas de pourcentage
    assert deltas["net_profit"] == Decimal("50.0")   # -20 -> -10 : amélioration


def test_trend_fills_empty_days(boutique):
    today = timezone.localdate()
    product = _product(boutique.compte, purchase_price=Decimal("300"))
    _stock(boutique, product, 20)
    _confirmed_sale(boutique, product, 1, 1000, sale_date=today - timedelta(days=2))
    data = reports.trend(boutique, today - timedelta(days=3), today)
    assert data["granularity"] == "day" and len(data["points"]) == 4
    revenues = [p["revenue"] for p in data["points"]]
    assert revenues == [0, 1000, 0, 0] and data["points"][1]["height"] == 100
    assert data["points"][1]["margin"] == Decimal("700")


def test_trend_groups_by_month_on_long_period(boutique):
    today = timezone.localdate()
    data = reports.trend(boutique, today - timedelta(days=200), today)
    assert data["granularity"] == "month" and 6 <= len(data["points"]) <= 8


def test_page_renders_with_comparison_and_trend(logged, boutique):
    product = _product(boutique.compte, purchase_price=Decimal("300"))
    _stock(boutique, product, 20)
    _confirmed_sale(boutique, product, 1, 1000)
    html = logged.get(reverse("sales:margin_report")).content.decode()
    assert "Panier moyen" in html and "Ventes par jour" in html and reverse("sales:margin_report_export") in html


def test_export_excel(logged, boutique):
    product = _product(boutique.compte, purchase_price=Decimal("300"))
    _stock(boutique, product, 20)
    _confirmed_sale(boutique, product, 2, 500)
    response = logged.get(reverse("sales:margin_report_export"))
    assert response.status_code == 200 and "spreadsheetml" in response["Content-Type"]
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Synthèse", "Produits"]
    synthese = {row[0].value: row[1].value for row in book["Synthèse"].iter_rows(min_row=4) if row[0].value}
    assert synthese["Ventes (HT)"] == 1000 and synthese["Marge brute"] == 400
    assert book["Produits"].cell(row=2, column=1).value == product.name


def test_export_forbidden_for_cashier(client, boutique):
    user = User.objects.create_user(email="c@test.com", password="x")
    Membership.objects.create(user=user, boutique=boutique, role=Membership.CAISSIER, is_active=True)
    client.force_login(user)
    assert client.get(reverse("sales:margin_report_export")).status_code in (302, 403)
