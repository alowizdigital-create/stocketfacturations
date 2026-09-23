"""Export de la liste des commandes (voir sales:commande_export_pdf /
sales:commande_export_excel) — un tableau simple, une ligne par commande,
pour un archivage ou un envoi à quelqu'un qui n'a pas accès à
l'application. Distinct de apps.sales.pdf, qui génère le PDF d'UN document
(facture/devis/commande) : ici c'est toujours une liste."""

import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BRAND_HEX = "6f4a8e"

COLUMNS = ("N°", "Client", "Téléphone", "Date", "Livraison", "Statut", "Total TTC")


def _rows(commandes):
    """Une ligne par commande, déjà formatée en texte — partagé par les
    deux formats pour qu'ils ne puissent pas diverger."""
    return [
        [
            str(commande.commande_seq) if commande.commande_seq is not None else commande.number,
            commande.client.name if commande.client else "—",
            commande.client.phone if commande.client and commande.client.phone else "—",
            commande.issue_date.strftime("%d/%m/%Y"),
            commande.get_delivery_status_display(),
            commande.get_status_display(),
            f"{commande.total_ttc:,.0f}".replace(",", " ") + f" {commande.currency}",
        ]
        for commande in commandes
    ]


def _title(boutique, delivery_label):
    return f"Commandes {delivery_label} — {boutique.name}"


def build_commandes_pdf(commandes, *, boutique, delivery_label):
    """PDF A4 paysage : simple tableau, une ligne par commande — pensé pour
    être imprimé ou joint à un message, pas pour remplacer l'écran."""
    styles = getSampleStyleSheet()
    style_title = ParagraphStyle(
        "ExportTitle", parent=styles["Normal"], fontSize=16, leading=20, fontName="Helvetica-Bold"
    )
    style_meta = ParagraphStyle("ExportMeta", parent=styles["Normal"], fontSize=9, leading=12, textColor=colors.grey)
    style_header = ParagraphStyle(
        "ExportHeaderCell", parent=styles["Normal"], fontSize=9, leading=11,
        fontName="Helvetica-Bold", textColor=colors.white,
    )
    style_cell = ParagraphStyle("ExportCell", parent=styles["Normal"], fontSize=9, leading=12)

    buffer = io.BytesIO()
    margin = 14 * mm
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=margin, bottomMargin=margin, leftMargin=margin, rightMargin=margin,
        title=_title(boutique, delivery_label),
    )

    story = [
        Paragraph(_title(boutique, delivery_label), style_title),
        Paragraph(f"Généré le {timezone.localdate():%d/%m/%Y} — {len(commandes)} commande(s)", style_meta),
        Spacer(1, 6 * mm),
    ]

    header_row = [Paragraph(label, style_header) for label in COLUMNS]
    table_rows = [header_row] + [[Paragraph(cell, style_cell) for cell in row] for row in _rows(commandes)]
    table = Table(table_rows, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BRAND_HEX}")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#dddddd")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8f7fa")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)

    if not commandes:
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("Aucune commande dans cette catégorie.", style_meta))

    doc.build(story)
    return buffer.getvalue()


def build_commandes_xlsx(commandes, *, boutique, delivery_label):
    """Classeur Excel : une feuille, une ligne d'en-tête figée, colonnes
    dimensionnées à la volée — pour retravailler la liste (tri, filtre,
    calculs) plutôt que la lire telle quelle."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Commandes"

    sheet.append([_title(boutique, delivery_label)])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([f"Généré le {timezone.localdate():%d/%m/%Y}"])
    sheet.append([])

    header_row_index = sheet.max_row + 1
    sheet.append(list(COLUMNS))
    header_fill = PatternFill("solid", fgColor=BRAND_HEX)
    for cell in sheet[header_row_index]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in _rows(commandes):
        sheet.append(row)

    sheet.freeze_panes = sheet.cell(row=header_row_index + 1, column=1)

    widths = [max(len(str(row[i])) for row in ([list(COLUMNS)] + _rows(commandes))) for i in range(len(COLUMNS))]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 40)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
